"""CSV / Excel export of an event's participant list.

Not listed in the original `services/` tree (only `participants.py` and
`statistics.py` were named), but export formatting is substantial enough
logic that it does not belong inline inside a handler.
"""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.database.models import Registration
from app.utils.helpers import STATUS_LABEL_RU, format_datetime

EXPORT_HEADERS = [
    "Telegram ID",
    "Username",
    "Имя",
    "Фамилия",
    "Статус",
    "Дата регистрации",
    "Дата изменения статуса",
    "Дата check-in",
]


def _row_for(registration: Registration) -> list[str]:
    user = registration.user
    return [
        str(user.telegram_id),
        f"@{user.username}" if user.username else "",
        user.first_name or "",
        user.last_name or "",
        STATUS_LABEL_RU[registration.status],
        format_datetime(registration.registered_at),
        format_datetime(registration.updated_at),
        format_datetime(registration.checked_in_at) if registration.checked_in_at else "",
    ]


def build_csv(registrations: list[Registration]) -> io.BytesIO:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    for registration in registrations:
        writer.writerow(_row_for(registration))
    # Excel-friendly BOM so Cyrillic text opens correctly on Windows.
    return io.BytesIO(("\ufeff" + buffer.getvalue()).encode("utf-8"))


def build_xlsx(registrations: list[Registration]) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Участники"
    sheet.append(EXPORT_HEADERS)
    for registration in registrations:
        sheet.append(_row_for(registration))

    for index, header in enumerate(EXPORT_HEADERS, start=1):
        column = get_column_letter(index)
        sheet.column_dimensions[column].width = max(len(header) + 2, 18)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
